import sys
import json
import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def natural_sort_key(s):
    if s is None:
        return []
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', str(s).strip())]

def clean_price(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    
    s = str(val).strip()
    if not s:
        return 0.0
    
    # Remove currency symbols and non-numeric except . , -
    s = re.sub(r'[^\d.,\-]', '', s)
    if not s:
        return 0.0
    
    # Handle comma vs dot
    if '.' in s and ',' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        s = s.replace(',', '.')
    
    try:
        return float(s)
    except (ValueError, TypeError):
        return 0.0

def find_header_map_bs(header_row):
    level_map = []
    for col_idx, cell in enumerate(header_row):
        # Ignore columns 0, 1, 2 (Codigo, Descripcion, Categoria)
        if col_idx < 3 or not cell:
            continue
        h = str(cell).upper().strip()
        
        # 1. Guayana / Depósito 2 (ID 3)
        if re.search(r'\b(GUAYANA|DEP[OÓ]SITO\s*2|DEP2|DEP\s*2|ALMACEN\s*LA\s*GUAYANA)\b', h):
            level_map.append((col_idx, 3))
        # 2. Concordia (ID 4)
        elif re.search(r'\b(CONCORDIA)\b', h):
            level_map.append((col_idx, 4))
        # 3. Caracas (ID 5)
        elif re.search(r'\b(CARACAS|CCS)\b', h):
            level_map.append((col_idx, 5))
        # 4. Valencia (ID 6)
        elif re.search(r'\b(VALENCIA|VAL)\b', h):
            level_map.append((col_idx, 6))
        # 5. Barinas (ID 7)
        elif re.search(r'\b(BARINAS|BNS)\b', h):
            level_map.append((col_idx, 7))
        # 6. Maracaibo (ID 8)
        elif re.search(r'\b(MARACAIBO|MCBO)\b', h):
            level_map.append((col_idx, 8))
        # 7. Merida (ID 9)
        elif re.search(r'\b(M[EÉ]RIDA)\b', h):
            level_map.append((col_idx, 9))
        # 8. Nacional (ID 10) e.g. "PRECIO NACIONAL (SC)"
        elif re.search(r'\b(NACIONAL|NC)\b', h):
            level_map.append((col_idx, 10))
        # 9. San Cristóbal / Depósito 1 (ID 2)
        elif re.search(r'\b(CRIST[OÓ]BAL|SAN\s*CRIST|SC|DEP1|DEP[OÓ]SITO\s*1)\b', h):
            level_map.append((col_idx, 2))
    
    if not level_map:
        level_map = [(3, 2), (4, 3), (5, 4), (6, 5), (7, 6), (8, 7), (9, 8), (10, 9), (11, 10)]
    return level_map

def find_header_map_divisas(header_row):
    level_map = []
    for col_idx, cell in enumerate(header_row):
        # Ignore columns 0, 1, 2 (Codigo, Descripcion, Categoria)
        if col_idx < 3 or not cell:
            continue
        h = str(cell).upper().strip()
        
        # 1. Caracas (ID 13)
        if re.search(r'\b(CARACAS|CCS)\b', h):
            level_map.append((col_idx, 13))
        # 2. Valencia (ID 14)
        elif re.search(r'\b(VALENCIA|VAL)\b', h):
            level_map.append((col_idx, 14))
        # 3. Barinas (ID 15)
        elif re.search(r'\b(BARINAS|BNS)\b', h):
            level_map.append((col_idx, 15))
        # 4. Maracaibo (ID 16)
        elif re.search(r'\b(MARACAIBO|MCBO)\b', h):
            level_map.append((col_idx, 16))
        # 5. Merida (ID 17)
        elif re.search(r'\b(M[EÉ]RIDA)\b', h):
            level_map.append((col_idx, 17))
        # 6. Concordia (ID 11)
        elif re.search(r'\b(CONCORDIA)\b', h):
            level_map.append((col_idx, 11))
        # 7. Nacional (ID 12) e.g. "DIVISA NC (DEP1)"
        elif re.search(r'\b(NACIONAL|NC)\b', h):
            level_map.append((col_idx, 12))
        # 8. San Cristóbal (ID 11) e.g. "DIVISA SC (DEP1)"
        elif re.search(r'\b(CRIST[OÓ]BAL|SAN\s*CRIST|SC|DEP1|DEP[OÓ]SITO\s*1)\b', h):
            level_map.append((col_idx, 11))
    
    if not level_map:
        level_map = [(3, 11), (4, 12), (5, 13), (6, 14), (7, 15), (8, 16), (9, 17)]
    return level_map

def parse_precios_bs(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_row = rows[0]
    level_map = find_header_map_bs(header_row)

    items = []
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        raw_code = row[0]
        if isinstance(raw_code, float) and raw_code.is_integer():
            code = str(int(raw_code)).strip()
        else:
            code = str(raw_code).strip()
        if code.endswith('.0') and code[:-2].isdigit():
            code = code[:-2]
        if not code:
            continue

        for col_idx, nivel_id in level_map:
            if col_idx < len(row):
                price = clean_price(row[col_idx])
                items.append({
                    "codigo": code,
                    "nivel_precio_id": nivel_id,
                    "moneda_id": 1,
                    "precio": price
                })

    items.sort(key=lambda item: natural_sort_key(item.get("codigo", "")))
    return items

def parse_precios_divisas(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_row = rows[0]
    level_map = find_header_map_divisas(header_row)

    items = []
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        raw_code = row[0]
        if isinstance(raw_code, float) and raw_code.is_integer():
            code = str(int(raw_code)).strip()
        else:
            code = str(raw_code).strip()
        if code.endswith('.0') and code[:-2].isdigit():
            code = code[:-2]
        if not code:
            continue

        for col_idx, nivel_id in level_map:
            if col_idx < len(row):
                price = clean_price(row[col_idx])
                items.append({
                    "codigo": code,
                    "nivel_precio_id": nivel_id,
                    "moneda_id": 2,
                    "precio": price
                })

    items.sort(key=lambda item: natural_sort_key(item.get("codigo", "")))
    return items

def write_excel_bs(json_data, output_path):
    # Sort data by codigo naturally (e.g. 01001, 01002, ...)
    sorted_data = sorted(json_data, key=lambda item: natural_sort_key(item.get("codigo", "")))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PRECIOS BS"

    headers = [
        "CÓDIGO", "DESCRIPCIÓN", "CATEGORÍA",
        "PRECIO DEPOSITO 1 SC", "PRECIO DEPOSITO 2 ALMACEN LA GUAYANA", "PRECIO CONCORDIA",
        "PRECIO CARACAS", "PRECIO VALENCIA", "PRECIO BARINAS",
        "PRECIO MARACAIBO", "PRECIO MERIDA", "PRECIO NACIONAL (SC)"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in sorted_data:
        ws.append([
            str(item.get("codigo", "")),
            str(item.get("nombre", "")),
            item.get("categoria_id", ""),
            float(item.get("sc", 0)),
            float(item.get("guayana", 0)),
            float(item.get("concordia", 0)),
            float(item.get("caracas", 0)),
            float(item.get("valencia", 0)),
            float(item.get("barinas", 0)),
            float(item.get("maracaibo", 0)),
            float(item.get("merida", 0)),
            float(item.get("nacional", 0)),
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 15)

    wb.save(output_path)

def write_excel_divisas(json_data, output_path):
    # Sort data by codigo naturally (e.g. 01001, 01002, ...)
    sorted_data = sorted(json_data, key=lambda item: natural_sort_key(item.get("codigo", "")))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PRECIOS DIVISAS"

    headers = [
        "CÓDIGO", "DESCRIPCIÓN", "CATEGORÍA",
        "DIVISA SC (DEP1)", "DIVISA NC (DEP1)", "DIVISA CARACAS",
        "DIVISA VALENCIA", "DIVISA BARINAS", "DIVISA MARACAIBO",
        "DIVISA MERIDA"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in sorted_data:
        ws.append([
            str(item.get("codigo", "")),
            str(item.get("nombre", "")),
            item.get("categoria_id", ""),
            float(item.get("sc", 0)),
            float(item.get("nacional", 0)),
            float(item.get("caracas", 0)),
            float(item.get("valencia", 0)),
            float(item.get("barinas", 0)),
            float(item.get("maracaibo", 0)),
            float(item.get("merida", 0)),
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 15)

    wb.save(output_path)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python procesar_excel_precios.py <parse-bs|parse-divisas|write-bs|write-divisas> <input/output path> [json_path]")
        sys.exit(1)

    cmd = sys.argv[1]
    path = sys.argv[2]

    if cmd == "parse-bs":
        res = parse_precios_bs(path)
        print(json.dumps(res))
    elif cmd == "parse-divisas":
        res = parse_precios_divisas(path)
        print(json.dumps(res))
    elif cmd == "write-bs":
        json_path = sys.argv[3]
        with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        write_excel_bs(data, path)
        print("OK")
    elif cmd == "write-divisas":
        json_path = sys.argv[3]
        with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        write_excel_divisas(data, path)
        print("OK")
